"""
Тесты рендера документов ОРД (раздел 8 ТЗ, п. 4.4 ТЗ, сессии B15, B16).

Этот файл:
- HTTP-слой рендера (GET /api/orders/{id}/render и
  GET /api/orders/{id}/render/status/{task_id});
- сервисный слой document_render_service: render_order_docx() и
  convert_docx_to_pdf().

Реальные Celery/Redis и LibreOffice в тестах НЕ используются:
- render_order_pdf (Celery-задача) и AsyncResult подменяются в app.api.orders
  через monkeypatch;
- subprocess.run (вызов soffice) подменяется фейком, который имитирует
  LibreOffice: кладёт source.pdf в --outdir либо возвращает ошибку;
- шаблон .docx для тестов рендера создаётся во временной папке прямо в тесте
  (python-docx входит в зависимости docxtpl), поэтому тесты не зависят от
  файлов в backend/app/templates/orders/ — кроме отдельного теста, который
  специально проверяет, что поставляемые шаблоны рендерятся.
"""
import base64
import os
import subprocess
import uuid
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.core.security import create_access_token, hash_password
from app.models.department import Department
from app.models.document_template import DocumentTemplate, DocumentTemplateType
from app.models.user import User, UserRole
from app.services import document_render_service
from app.services.document_render_service import (
    DocumentConversionError,
    TemplateFileNotFoundError,
    convert_docx_to_pdf,
    render_order_docx,
)

DOCX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _auth_headers(user: User) -> dict[str, str]:
    token = create_access_token(user.id, user.role.value)
    return {"Authorization": f"Bearer {token}"}


def _make_user(db_session: Session, department: Department, role: UserRole) -> User:
    user = User(
        full_name=f"Тестовый {role.value}",
        department_id=department.id,
        role=role,
        login=f"test_{role.name.lower()}_{uuid.uuid4().hex[:8]}",
        password_hash=hash_password("TestPassword123!"),
        is_active=True,
    )
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return user


def _make_template(db_session: Session, file_path: Path | str) -> DocumentTemplate:
    template = DocumentTemplate(
        name="Тестовый шаблон рендера",
        type=DocumentTemplateType.WORK_ORDER,
        file_path=str(file_path),
        field_schema=[],
        min_approver_role=UserRole.IT_HEAD,
    )
    db_session.add(template)
    db_session.commit()
    db_session.refresh(template)
    return template


def _make_docx_template_file(tmp_path: Path) -> Path:
    """Минимальный .docx с одним Jinja2-плейсхолдером {{ field }} —
    ровно то, что подставляет docxtpl из Order.fields."""
    document = Document()
    document.add_paragraph("Значение поля: {{ field }}")
    path = tmp_path / "template.docx"
    document.save(str(path))
    return path


def _create_order(client: TestClient, author: User, template: DocumentTemplate) -> dict:
    response = client.post(
        "/api/orders",
        json={
            "type": template.type.value,
            "template_id": str(template.id),
            "fields": {"field": "значение"},
        },
        headers=_auth_headers(author),
    )
    assert response.status_code == 201
    return response.json()


def _patch_async_result(
    monkeypatch: pytest.MonkeyPatch, *, state: str, result: object = None
) -> list[str]:
    """Подменяет AsyncResult в app.api.orders: возвращает объект с заданными
    state/result. Список запрошенных task_id возвращается для assert'ов."""
    requested: list[str] = []

    def _factory(task_id: str, app: object = None) -> SimpleNamespace:
        requested.append(task_id)
        return SimpleNamespace(state=state, result=result)

    monkeypatch.setattr("app.api.orders.AsyncResult", _factory)
    return requested


# --- GET /api/orders/{id}/render?format=docx ---


def test_render_docx_returns_file_with_substituted_fields(
    client: TestClient, db_session: Session, engineer_user: User, tmp_path: Path
) -> None:
    template = _make_template(db_session, _make_docx_template_file(tmp_path))
    order = _create_order(client, engineer_user, template)

    response = client.get(
        f"/api/orders/{order['id']}/render?format=docx", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_MEDIA_TYPE
    assert f'filename="{order["id"]}.docx"' in response.headers["content-disposition"]
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert "значение" in text
    assert "{{" not in text


@pytest.mark.parametrize(
    "role", [UserRole.ENGINEER, UserRole.IT_HEAD, UserRole.EXECUTIVE, UserRole.ADMIN]
)
def test_render_docx_allowed_for_engineer_and_above(
    client: TestClient,
    db_session: Session,
    department: Department,
    engineer_user: User,
    tmp_path: Path,
    role: UserRole,
) -> None:
    """Раздел 4.4 ТЗ: рендер доступен Engineer+; Executive включён явно
    (докстринг api/orders.py) — он должен видеть документ, который согласует."""
    template = _make_template(db_session, _make_docx_template_file(tmp_path))
    order = _create_order(client, engineer_user, template)
    viewer = _make_user(db_session, department, role)

    response = client.get(
        f"/api/orders/{order['id']}/render?format=docx", headers=_auth_headers(viewer)
    )

    assert response.status_code == 200


def test_render_docx_returns_500_when_template_file_is_missing(
    client: TestClient, db_session: Session, engineer_user: User, tmp_path: Path
) -> None:
    """Ошибка конфигурации/деплоя (шаблон не попал в образ) — 500, а не 404."""
    template = _make_template(db_session, tmp_path / "does_not_exist.docx")
    order = _create_order(client, engineer_user, template)

    response = client.get(
        f"/api/orders/{order['id']}/render?format=docx", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 500


# --- GET /api/orders/{id}/render?format=pdf ---


def test_render_pdf_enqueues_task_and_returns_202(
    client: TestClient,
    db_session: Session,
    engineer_user: User,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[str] = []

    class _FakeTask:
        @staticmethod
        def delay(order_id: str) -> SimpleNamespace:
            calls.append(order_id)
            return SimpleNamespace(id="fake-task-id")

    monkeypatch.setattr("app.api.orders.render_order_pdf", _FakeTask)
    template = _make_template(db_session, tmp_path / "unused.docx")
    order = _create_order(client, engineer_user, template)

    response = client.get(
        f"/api/orders/{order['id']}/render?format=pdf", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 202
    body = response.json()
    assert body["task_id"] == "fake-task-id"
    assert body["status_url"] == f"/api/orders/{order['id']}/render/status/fake-task-id"
    assert calls == [order["id"]]


# --- общие проверки /render ---


def test_render_requires_format_parameter(
    client: TestClient, db_session: Session, engineer_user: User, tmp_path: Path
) -> None:
    template = _make_template(db_session, tmp_path / "unused.docx")
    order = _create_order(client, engineer_user, template)

    response = client.get(f"/api/orders/{order['id']}/render", headers=_auth_headers(engineer_user))

    assert response.status_code == 422


def test_render_rejects_unknown_format(
    client: TestClient, db_session: Session, engineer_user: User, tmp_path: Path
) -> None:
    template = _make_template(db_session, tmp_path / "unused.docx")
    order = _create_order(client, engineer_user, template)

    response = client.get(
        f"/api/orders/{order['id']}/render?format=xlsx", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 422


def test_render_not_found_for_unknown_order(client: TestClient, engineer_user: User) -> None:
    response = client.get(
        f"/api/orders/{uuid.uuid4()}/render?format=docx", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 404


def test_render_forbidden_for_user_role(
    client: TestClient, db_session: Session, department: Department
) -> None:
    shadow_user = _make_user(db_session, department, UserRole.USER)

    response = client.get(
        f"/api/orders/{uuid.uuid4()}/render?format=docx", headers=_auth_headers(shadow_user)
    )

    assert response.status_code == 403


# --- GET /api/orders/{id}/render/status/{task_id} ---


@pytest.mark.parametrize("celery_state", ["PENDING", "STARTED", "RETRY"])
def test_render_status_pending_states_return_202(
    client: TestClient,
    db_session: Session,
    engineer_user: User,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    celery_state: str,
) -> None:
    requested = _patch_async_result(monkeypatch, state=celery_state)
    template = _make_template(db_session, tmp_path / "unused.docx")
    order = _create_order(client, engineer_user, template)

    response = client.get(
        f"/api/orders/{order['id']}/render/status/task-123", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 202
    assert response.json() == {"status": "pending"}
    assert requested == ["task-123"]


def test_render_status_success_returns_pdf_bytes(
    client: TestClient,
    db_session: Session,
    engineer_user: User,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pdf_bytes = b"%PDF-1.4 fake pdf content"
    template = _make_template(db_session, tmp_path / "unused.docx")
    order = _create_order(client, engineer_user, template)
    _patch_async_result(
        monkeypatch,
        state="SUCCESS",
        result={
            "filename": f"{order['id']}.pdf",
            "content_b64": base64.b64encode(pdf_bytes).decode("ascii"),
        },
    )

    response = client.get(
        f"/api/orders/{order['id']}/render/status/task-123", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert f'filename="{order["id"]}.pdf"' in response.headers["content-disposition"]
    assert response.content == pdf_bytes


def test_render_status_failure_returns_500_with_reason(
    client: TestClient,
    db_session: Session,
    engineer_user: User,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _patch_async_result(monkeypatch, state="FAILURE", result=RuntimeError("soffice упал"))
    template = _make_template(db_session, tmp_path / "unused.docx")
    order = _create_order(client, engineer_user, template)

    response = client.get(
        f"/api/orders/{order['id']}/render/status/task-123", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 500
    assert "soffice упал" in response.json()["detail"]


def test_render_status_unknown_intermediate_state_is_treated_as_not_ready(
    client: TestClient,
    db_session: Session,
    engineer_user: User,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Промежуточное состояние Celery вне списка PENDING/STARTED/RETRY
    (например, RECEIVED) — 202 с состоянием в нижнем регистре, а не ошибка."""
    _patch_async_result(monkeypatch, state="RECEIVED")
    template = _make_template(db_session, tmp_path / "unused.docx")
    order = _create_order(client, engineer_user, template)

    response = client.get(
        f"/api/orders/{order['id']}/render/status/task-123", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 202
    assert response.json() == {"status": "received"}


def test_render_status_not_found_for_unknown_order(
    client: TestClient, engineer_user: User, monkeypatch: pytest.MonkeyPatch
) -> None:
    requested = _patch_async_result(monkeypatch, state="PENDING")

    response = client.get(
        f"/api/orders/{uuid.uuid4()}/render/status/task-123", headers=_auth_headers(engineer_user)
    )

    assert response.status_code == 404
    assert requested == []  # до Celery дело не дошло: заказ проверяется первым


def test_render_status_forbidden_for_user_role(
    client: TestClient, db_session: Session, department: Department
) -> None:
    shadow_user = _make_user(db_session, department, UserRole.USER)

    response = client.get(
        f"/api/orders/{uuid.uuid4()}/render/status/task-123", headers=_auth_headers(shadow_user)
    )

    assert response.status_code == 403


# --- B16: сервисный слой document_render_service ---


def _make_order_stub(template_path: Path, fields: dict) -> SimpleNamespace:
    """Минимальная замена Order для render_order_docx(): сервису нужны только
    order.template.file_path, order.fields и order.template_id (для текста
    ошибки). БД не требуется."""
    return SimpleNamespace(
        id=uuid.uuid4(),
        template_id=uuid.uuid4(),
        template=SimpleNamespace(file_path=str(template_path)),
        fields=fields,
    )


def test_render_order_docx_substitutes_fields(tmp_path: Path) -> None:
    order = _make_order_stub(_make_docx_template_file(tmp_path), {"field": "значение"})

    result = render_order_docx(order)

    assert isinstance(result, bytes)
    text = "\n".join(p.text for p in Document(BytesIO(result)).paragraphs)
    assert "Значение поля: значение" in text
    assert "{{" not in text


def test_render_order_docx_raises_when_template_file_is_missing(tmp_path: Path) -> None:
    missing_path = tmp_path / "does_not_exist.docx"
    order = _make_order_stub(missing_path, {"field": "значение"})

    with pytest.raises(TemplateFileNotFoundError) as exc_info:
        render_order_docx(order)

    assert str(missing_path) in str(exc_info.value)


@pytest.mark.parametrize("template_name", ["purchase_request.docx", "write_off_act.docx"])
def test_shipped_templates_render_without_errors(template_name: str) -> None:
    """Поставляемые шаблоны ОРД (B15) должны быть валидными .docx с корректным
    Jinja2-синтаксисом: рендер с пустыми fields не должен падать. Ловит
    поломку шаблона при ручной правке в Word, а не у пользователя в проде."""
    templates_dir = Path(document_render_service.__file__).resolve().parent.parent / "templates" / "orders"
    order = _make_order_stub(templates_dir / template_name, {})

    result = render_order_docx(order)

    assert Document(BytesIO(result)).paragraphs is not None


def _install_fake_soffice(
    monkeypatch: pytest.MonkeyPatch,
    *,
    returncode: int = 0,
    write_pdf: bool = True,
    stderr: bytes = b"",
    pdf_bytes: bytes = b"%PDF-1.4 fake pdf",
) -> list[dict]:
    """Подменяет subprocess.run фейком LibreOffice. Фейк читает source.docx,
    при write_pdf=True кладёт source.pdf в --outdir и возвращает заданный
    returncode/stderr. Каждый вызов записывается в возвращаемый список
    (аргументы, kwargs, содержимое docx, пути временных каталогов)."""
    calls: list[dict] = []
    profile_prefix = "-env:UserInstallation=file://"

    def _fake_run(cmd: list[str], **kwargs: object) -> SimpleNamespace:
        outdir = Path(cmd[cmd.index("--outdir") + 1])
        profile_arg = next(arg for arg in cmd if arg.startswith(profile_prefix))
        calls.append(
            {
                "cmd": cmd,
                "kwargs": kwargs,
                "outdir": outdir,
                "profile_dir": Path(profile_arg.removeprefix(profile_prefix)),
                "docx_content": Path(cmd[-1]).read_bytes(),
            }
        )
        if write_pdf:
            (outdir / "source.pdf").write_bytes(pdf_bytes)
        return SimpleNamespace(returncode=returncode, stderr=stderr)

    monkeypatch.setattr(document_render_service.subprocess, "run", _fake_run)
    return calls


def test_convert_docx_to_pdf_success_uses_expected_soffice_invocation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Готчи LibreOffice из B15 (ways-of-working): SAL_USE_VCLPLUGIN=svp,
    свой -env:UserInstallation на вызов, --norestore, таймаут."""
    pdf_bytes = "%PDF-1.4 результат".encode()
    calls = _install_fake_soffice(monkeypatch, pdf_bytes=pdf_bytes)
    docx_bytes = b"PK fake docx content"

    result = convert_docx_to_pdf(docx_bytes)

    assert result == pdf_bytes
    assert len(calls) == 1
    call = calls[0]
    cmd = call["cmd"]
    assert cmd[0] == "soffice"
    assert "--headless" in cmd
    assert "--norestore" in cmd
    assert cmd[cmd.index("--convert-to") + 1] == "pdf"
    assert cmd[-1].endswith("source.docx")
    assert call["docx_content"] == docx_bytes

    kwargs = call["kwargs"]
    assert kwargs["timeout"] == document_render_service._LIBREOFFICE_TIMEOUT_SECONDS
    assert kwargs["capture_output"] is True
    assert kwargs["check"] is False
    assert kwargs["env"]["SAL_USE_VCLPLUGIN"] == "svp"
    assert kwargs["env"]["PATH"] == os.environ["PATH"]  # окружение не потеряно, soffice находится

    # Временные каталоги (рабочий и профиль LibreOffice) убраны после вызова
    assert not call["outdir"].exists()
    assert not call["profile_dir"].exists()


def test_convert_docx_to_pdf_uses_unique_profile_per_call(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Общий профиль между вызовами приводит к lock-contention (B15) —
    у каждого вызова должен быть свой UserInstallation."""
    calls = _install_fake_soffice(monkeypatch)

    convert_docx_to_pdf(b"PK first")
    convert_docx_to_pdf(b"PK second")

    assert len(calls) == 2
    assert calls[0]["profile_dir"] != calls[1]["profile_dir"]


@pytest.mark.parametrize(
    ("returncode", "write_pdf"),
    [
        (1, False),  # soffice упал и PDF нет
        (1, True),  # код ошибки важнее случайно оказавшегося файла
        (0, False),  # код 0, но PDF не создан
    ],
)
def test_convert_docx_to_pdf_raises_on_failure(
    monkeypatch: pytest.MonkeyPatch, returncode: int, write_pdf: bool
) -> None:
    _install_fake_soffice(
        monkeypatch,
        returncode=returncode,
        write_pdf=write_pdf,
        stderr="ошибка конвертации".encode(),
    )

    with pytest.raises(DocumentConversionError) as exc_info:
        convert_docx_to_pdf(b"PK fake docx")

    message = str(exc_info.value)
    assert f"кодом {returncode}" in message
    assert "ошибка конвертации" in message


def test_convert_docx_to_pdf_raises_on_timeout(monkeypatch: pytest.MonkeyPatch) -> None:
    def _raise_timeout(cmd: list[str], **kwargs: object) -> None:
        raise subprocess.TimeoutExpired(cmd, kwargs["timeout"])

    monkeypatch.setattr(document_render_service.subprocess, "run", _raise_timeout)

    with pytest.raises(DocumentConversionError) as exc_info:
        convert_docx_to_pdf(b"PK fake docx")

    assert str(document_render_service._LIBREOFFICE_TIMEOUT_SECONDS) in str(exc_info.value)
    assert isinstance(exc_info.value.__cause__, subprocess.TimeoutExpired)


def test_convert_docx_to_pdf_tolerates_non_utf8_stderr(monkeypatch: pytest.MonkeyPatch) -> None:
    """stderr декодируется с errors="replace": мусор в выводе soffice не должен
    превращать DocumentConversionError в UnicodeDecodeError."""
    _install_fake_soffice(monkeypatch, returncode=1, write_pdf=False, stderr=b"\xff\xfe broken")

    with pytest.raises(DocumentConversionError):
        convert_docx_to_pdf(b"PK fake docx")
